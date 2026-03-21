from __future__ import annotations

import re
from copy import copy
from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.worksheet import Worksheet

from config import (
    ARCHIVE_MARKER,
    COLUMN_DEFINITIONS,
    COLUMNS_BY_KEY,
    DATA_START_ROW,
    DEFAULT_SELECT_OPTIONS,
    FIELD_KEYS,
    NUMBER_PREFIX,
    NUMBER_SUFFIX_DEFAULT,
)

NUMBER_REGEX = re.compile(rf"^{NUMBER_PREFIX}_(?P<type>[A-Za-z0-9]+)_(?P<sequence>\d{{3}})_(?P<revision>\d{{2}})$")


class ExcelHandlerError(Exception):
    """Raised when workbook operations fail."""


class ExcelHandler:
    def __init__(self, workbook_path: Path):
        self.workbook_path = Path(workbook_path)

    def _load_workbook(self):
        if not self.workbook_path.exists():
            raise ExcelHandlerError(f"Fichier introuvable : {self.workbook_path}")
        try:
            return load_workbook(self.workbook_path)
        except PermissionError as exc:
            raise ExcelHandlerError("Le fichier Excel semble verrouillé (PermissionError).") from exc

    def _get_sheet(self, workbook, sheet_name: str) -> Worksheet:
        if sheet_name not in workbook.sheetnames:
            raise ExcelHandlerError(f"Onglet introuvable : {sheet_name}")
        return workbook[sheet_name]

    @staticmethod
    def _last_data_row(sheet: Worksheet) -> int:
        for row in range(sheet.max_row, DATA_START_ROW - 1, -1):
            if any(sheet[f"{cfg.letter}{row}"].value not in (None, "") for cfg in COLUMN_DEFINITIONS):
                return row
        return DATA_START_ROW - 1

    @staticmethod
    def _next_index(sheet: Worksheet) -> int:
        max_index = 0
        for row in range(DATA_START_ROW, sheet.max_row + 1):
            current = sheet[f"A{row}"].value
            try:
                max_index = max(max_index, int(current))
            except (TypeError, ValueError):
                continue
        return max_index + 1 if max_index else 1

    @staticmethod
    def _normalize_type(type_value: str) -> str:
        normalized = re.sub(r"[^A-Za-z0-9]+", "", type_value.upper())
        if not normalized:
            raise ExcelHandlerError("Le TYPE d'identification est obligatoire pour générer le NUMBER.")
        return normalized

    @classmethod
    def _build_number(cls, type_value: str, next_index: int, revision: int = NUMBER_SUFFIX_DEFAULT) -> str:
        normalized_type = cls._normalize_type(type_value)
        return f"{NUMBER_PREFIX}_{normalized_type}_{next_index:03d}_{revision:02d}"

    @classmethod
    def _increment_number_revision(cls, current_number: str, type_value: str) -> str:
        match = NUMBER_REGEX.match(str(current_number or "").strip())
        if match:
            original_type = match.group("type")
            sequence = int(match.group("sequence"))
            revision = int(match.group("revision")) + 1
            return cls._build_number(original_type, sequence, revision)
        return cls._build_number(type_value, 1, 1)

    def preview_next_number(self, sheet_name: str, type_value: str) -> str:
        wb = self._load_workbook()
        sheet = self._get_sheet(wb, sheet_name)
        next_index = self._next_index(sheet)
        number = self._build_number(type_value, next_index)
        wb.close()
        return number

    def preview_updated_number(self, current_number: str, type_value: str) -> str:
        return self._increment_number_revision(current_number, type_value)

    def list_sheets(self) -> List[str]:
        wb = self._load_workbook()
        names = wb.sheetnames
        wb.close()
        return names

    def get_dynamic_options(self, sheet_name: str) -> Dict[str, List[str]]:
        wb = self._load_workbook()
        sheet = self._get_sheet(wb, sheet_name)
        options = {key: values[:] for key, values in DEFAULT_SELECT_OPTIONS.items()}

        last_row = self._last_data_row(sheet)
        for key in [cfg.key for cfg in COLUMN_DEFINITIONS if cfg.dropdown]:
            col = COLUMNS_BY_KEY[key].letter
            seen = set(options.get(key, []))
            dynamic = []
            for row in range(DATA_START_ROW, last_row + 1):
                val = sheet[f"{col}{row}"].value
                if val in (None, ""):
                    continue
                val_str = str(val)
                if val_str not in seen:
                    seen.add(val_str)
                    dynamic.append(val_str)
            options[key] = options.get(key, []) + sorted(dynamic)
        wb.close()
        return options

    @staticmethod
    def _copy_row_style_and_formulas(sheet: Worksheet, source_row: int, target_row: int) -> None:
        for cfg in COLUMN_DEFINITIONS:
            source = sheet[f"{cfg.letter}{source_row}"]
            target = sheet[f"{cfg.letter}{target_row}"]
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            if source.protection:
                target.protection = copy(source.protection)
            if source.alignment:
                target.alignment = copy(source.alignment)
            if isinstance(source.value, str) and source.value.startswith("="):
                target.value = Translator(source.value, origin=f"{cfg.letter}{source_row}").translate_formula(
                    f"{cfg.letter}{target_row}"
                )

    @staticmethod
    def _apply_payload(sheet: Worksheet, row: int, payload: Dict[str, str]) -> None:
        for key in FIELD_KEYS:
            col = COLUMNS_BY_KEY[key].letter
            sheet[f"{col}{row}"] = payload.get(key, "")

    def add_requirement(self, sheet_name: str, payload: Dict[str, str]) -> int:
        wb = self._load_workbook()
        sheet = self._get_sheet(wb, sheet_name)
        insert_row = DATA_START_ROW
        next_index = self._next_index(sheet)
        payload_with_number = dict(payload)
        payload_with_number["number"] = self._build_number(payload_with_number.get("type", ""), next_index)

        sheet.insert_rows(insert_row, amount=1)
        if sheet.max_row >= insert_row + 1:
            self._copy_row_style_and_formulas(sheet, insert_row + 1, insert_row)

        sheet[f"A{insert_row}"] = next_index
        self._apply_payload(sheet, insert_row, payload_with_number)

        wb.save(self.workbook_path)
        wb.close()
        return insert_row

    def list_requirements(
        self,
        sheet_name: str,
        keyword: str = "",
        type_filter: str = "",
        applicability_filter: str = "",
        include_archived: bool = True,
    ) -> List[Dict[str, str]]:
        wb = self._load_workbook()
        sheet = self._get_sheet(wb, sheet_name)
        last_row = self._last_data_row(sheet)
        keyword_lower = keyword.lower().strip()
        type_filter_upper = type_filter.upper().strip()
        applicability_filter_upper = applicability_filter.upper().strip()
        results: List[Dict[str, str]] = []

        for row in range(DATA_START_ROW, last_row + 1):
            number = str(sheet[f"B{row}"].value or "")
            french_name = str(sheet[f"C{row}"].value or "")
            english_name = str(sheet[f"D{row}"].value or "")
            type_value = str(sheet[f"E{row}"].value or "")
            applicability = str(sheet[f"M{row}"].value or "")
            archived = french_name.startswith(ARCHIVE_MARKER)

            haystack = " ".join([number, french_name, english_name, type_value, applicability]).lower()
            if keyword_lower and keyword_lower not in haystack:
                continue
            if type_filter_upper and type_value.upper() != type_filter_upper:
                continue
            if applicability_filter_upper and applicability.upper() != applicability_filter_upper:
                continue
            if not include_archived and archived:
                continue

            results.append(
                {
                    "row": row,
                    "number": number,
                    "french_name": french_name,
                    "english_name": english_name,
                    "type": type_value,
                    "applicability": applicability,
                    "archived": archived,
                }
            )
        wb.close()
        return results

    def search_requirements(self, sheet_name: str, keyword: str) -> List[Dict[str, str]]:
        return self.list_requirements(sheet_name=sheet_name, keyword=keyword)

    def load_requirement(self, sheet_name: str, row: int) -> Dict[str, str]:
        wb = self._load_workbook()
        sheet = self._get_sheet(wb, sheet_name)
        data = {key: str(sheet[f"{COLUMNS_BY_KEY[key].letter}{row}"].value or "") for key in FIELD_KEYS}
        wb.close()
        return data

    def update_requirement(self, sheet_name: str, row: int, payload: Dict[str, str]) -> str:
        wb = self._load_workbook()
        sheet = self._get_sheet(wb, sheet_name)
        payload_with_number = dict(payload)
        current_number = str(sheet[f"B{row}"].value or payload_with_number.get("number", ""))
        payload_with_number["number"] = self._increment_number_revision(current_number, payload_with_number.get("type", ""))
        self._apply_payload(sheet, row, payload_with_number)
        wb.save(self.workbook_path)
        wb.close()
        return payload_with_number["number"]

    def archive_requirement(self, sheet_name: str, row: int) -> None:
        wb = self._load_workbook()
        sheet = self._get_sheet(wb, sheet_name)
        current = str(sheet[f"C{row}"].value or "")
        if not current.startswith(ARCHIVE_MARKER):
            sheet[f"C{row}"] = f"{ARCHIVE_MARKER} {current}".strip()
        wb.save(self.workbook_path)
        wb.close()
